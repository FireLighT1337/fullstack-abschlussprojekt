"""Document loading and text extraction for RAG.

Supports: PDF, Excel, PowerPoint, Word, plain text.
Images (jpg/jpeg) are saved but not processed for RAG in this MVP.
"""

import os
from pathlib import Path
from typing import List, Optional

# LangChain loaders
from langchain_community.document_loaders import (
    PyPDFLoader,
    UnstructuredExcelLoader,
    UnstructuredPowerPointLoader,
    TextLoader,
)
from langchain_core.documents import Document


DATA_DIR = Path(os.environ.get("DATA_DIR", "/tmp/bernd_data"))
DATA_DIR.mkdir(parents=True, exist_ok=True)


def save_upload(file_name: str, file_bytes: bytes) -> Path:
    """Save uploaded file to disk and return the path."""
    path = DATA_DIR / file_name
    # avoid collisions
    counter = 1
    stem = path.stem
    suffix = path.suffix
    while path.exists():
        path = DATA_DIR / f"{stem}_{counter}{suffix}"
        counter += 1
    path.write_bytes(file_bytes)
    return path


def detect_mime_from_name(name: str) -> str:
    name_lower = name.lower()
    if name_lower.endswith(".pdf"):
        return "application/pdf"
    if name_lower.endswith((".xls", ".xlsx")):
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if name_lower.endswith((".ppt", ".pptx")):
        return "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    if name_lower.endswith((".doc", ".docx")):
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if name_lower.endswith((".txt", ".md", ".csv", ".json")):
        return "text/plain"
    if name_lower.endswith((".jpg", ".jpeg", ".png", ".gif")):
        return "image/*"
    return "application/octet-stream"


def load_documents(file_path: Path, mime_type: str) -> List[Document]:
    """Extract text from a file using the appropriate LangChain loader."""
    str_path = str(file_path)
    name_lower = file_path.name.lower()

    try:
        if name_lower.endswith(".pdf"):
            loader = PyPDFLoader(str_path)
            return loader.load()

        if name_lower.endswith((".xls", ".xlsx")):
            # UnstructuredExcelLoader needs the file path
            loader = UnstructuredExcelLoader(str_path, mode="elements")
            return loader.load()

        if name_lower.endswith((".ppt", ".pptx")):
            loader = UnstructuredPowerPointLoader(str_path)
            return loader.load()

        if name_lower.endswith((".doc", ".docx")):
            # Try unstructured if available, otherwise fallback to simple text extraction
            try:
                from langchain_community.document_loaders import UnstructuredWordDocumentLoader
                loader = UnstructuredWordDocumentLoader(str_path)
                return loader.load()
            except Exception:
                # Fallback: try python-docx directly
                try:
                    from docx import Document as DocxDocument
                    doc = DocxDocument(str_path)
                    text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
                    if text.strip():
                        return [Document(page_content=text, metadata={"source": str_path})]
                    return []
                except Exception:
                    return []

        if name_lower.endswith((".txt", ".md", ".csv", ".json", ".py", ".js", ".html", ".css")):
            loader = TextLoader(str_path, encoding="utf-8")
            return loader.load()

        # Images: not processed for text in MVP
        if name_lower.endswith((".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp")):
            return []

    except Exception as e:
        print(f"[Document Load Error] {file_path}: {e}")
        return []

    return []